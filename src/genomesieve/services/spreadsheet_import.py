import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ASSEMBLY_ACCESSION_PATTERN = re.compile(
    r"^(?:GCF|GCA)_\d+\.\d+$",
    re.IGNORECASE,
)


class SpreadsheetImportError(Exception):
    """Raised when a spreadsheet cannot be analyzed or imported."""

@dataclass(frozen=True)
class SpreadsheetColumn:
    """
    A column available for manual selection in the spreadsheet
    import interface.
    """

    column_index: int
    column_letter: str
    label: str

    examples: tuple[str, ...]

@dataclass(frozen=True)
class ColumnCandidate:
    """
    A spreadsheet column that appears to contain NCBI Assembly
    Accessions.
    """

    column_index: int
    column_letter: str
    label: str

    valid_accessions: int
    non_empty_values: int

    match_ratio: float
    score: float

    examples: tuple[str, ...]


@dataclass(frozen=True)
class SheetAnalysis:
    """
    Result of Assembly Accession column detection for one worksheet.
    """

    sheet_name: str

    available_columns: tuple[SpreadsheetColumn, ...]
    candidate_columns: tuple[ColumnCandidate, ...]

    detected_column_index: int | None
    detected_column_label: str | None

    ambiguous: bool


@dataclass(frozen=True)
class SpreadsheetAnalysis:
    """
    General structure detected in an imported spreadsheet.
    """

    file_path: Path
    sheets: tuple[SheetAnalysis, ...]


@dataclass(frozen=True)
class ImportedAccession:
    """
    One occurrence of an Assembly Accession in the spreadsheet.

    Duplicates are intentionally preserved.
    """

    accession: str

    sheet_name: str
    row_number: int

    column_index: int
    column_label: str


@dataclass(frozen=True)
class InvalidAccessionEntry:
    """
    A non-empty value found in the selected accession column that does
    not have the expected GCF_/GCA_ syntax.
    """

    value: str

    sheet_name: str
    row_number: int

    column_index: int
    column_label: str


@dataclass
class SpreadsheetImportResult:
    """
    Result of extracting accessions from the detected or manually
    selected columns.
    """

    analysis: SpreadsheetAnalysis

    entries: list[ImportedAccession]
    invalid_entries: list[InvalidAccessionEntry]

    @property
    def total_entries(self):
        return len(self.entries)

    @property
    def unique_accessions(self):
        return len(
            {
                entry.accession
                for entry in self.entries
            }
        )

    @property
    def duplicate_entries(self):
        return (
            self.total_entries
            - self.unique_accessions
        )

    @property
    def total_invalid_entries(self):
        return len(
            self.invalid_entries
        )


def is_valid_assembly_accession(value):
    """
    Return True if the value has the expected NCBI Assembly Accession
    syntax.

    Examples:
        GCF_022592395.1
        GCA_000276625.1
    """

    if value is None:
        return False

    value = str(value).strip()

    return bool(
        ASSEMBLY_ACCESSION_PATTERN.fullmatch(
            value
        )
    )


def analyze_spreadsheet(file_path):
    """
    Analyze all worksheets and automatically detect columns that appear
    to contain NCBI Assembly Accessions.

    No accessions are removed or deduplicated here.
    """

    path = _validate_file(file_path)

    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )

    except Exception as exc:
        raise SpreadsheetImportError(
            f"Could not open spreadsheet: {exc}"
        ) from exc

    try:
        sheet_analyses = []

        for worksheet in workbook.worksheets:

            analysis = _analyze_sheet(
                worksheet
            )

            sheet_analyses.append(
                analysis
            )

        return SpreadsheetAnalysis(
            file_path=path,
            sheets=tuple(
                sheet_analyses
            ),
        )

    finally:
        workbook.close()


def import_accessions(
    file_path,
    column_selection=None,
):
    """
    Extract Assembly Accessions from the spreadsheet.

    column_selection is optional and will later be supplied by the GUI.

    Example:

        {
            "Staphylococcus": 4,
            "Bacillus": 3,
        }

    The key is the worksheet name and the value is the 1-based column
    index.

    When a worksheet is not explicitly provided in column_selection,
    its automatically detected column is used.

    Duplicates are intentionally preserved.
    """

    analysis = analyze_spreadsheet(
        file_path
    )

    path = analysis.file_path

    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )

    except Exception as exc:
        raise SpreadsheetImportError(
            f"Could not open spreadsheet: {exc}"
        ) from exc

    entries = []
    invalid_entries = []

    try:
        analysis_by_sheet = {
            sheet.sheet_name: sheet
            for sheet in analysis.sheets
        }

        for worksheet in workbook.worksheets:

            sheet_analysis = (
                analysis_by_sheet[
                    worksheet.title
                ]
            )

            selected_column = None

            if (
                column_selection
                and worksheet.title
                in column_selection
            ):
                selected_column = (
                    column_selection[
                        worksheet.title
                    ]
                )

            else:
                selected_column = (
                    sheet_analysis
                    .detected_column_index
                )

            # No reliable accession column was found for this sheet.
            if selected_column is None:
                continue

            column_label = _get_candidate_label(
                sheet_analysis,
                selected_column,
            )

            first_accession_found = False

            for row_number, row in enumerate(
                worksheet.iter_rows(
                    values_only=True
                ),
                start=1,
            ):

                if (
                    selected_column
                    > len(row)
                ):
                    continue

                value = row[
                    selected_column - 1
                ]

                text = _normalize_value(
                    value
                )

                if not text:
                    continue

                if is_valid_assembly_accession(
                    text
                ):
                    first_accession_found = True

                    entries.append(
                        ImportedAccession(
                            accession=(
                                text.upper()
                            ),
                            sheet_name=(
                                worksheet.title
                            ),
                            row_number=(
                                row_number
                            ),
                            column_index=(
                                selected_column
                            ),
                            column_label=(
                                column_label
                            ),
                        )
                    )

                    continue

                # Ignore text before the first accession because it is
                # normally the header or introductory spreadsheet text.
                if not first_accession_found:
                    continue

                # Repeated headers are common in real laboratory
                # spreadsheets. Do not classify them as invalid data.
                if _looks_like_accession_header(
                    text
                ):
                    continue

                invalid_entries.append(
                    InvalidAccessionEntry(
                        value=text,
                        sheet_name=(
                            worksheet.title
                        ),
                        row_number=(
                            row_number
                        ),
                        column_index=(
                            selected_column
                        ),
                        column_label=(
                            column_label
                        ),
                    )
                )

        return SpreadsheetImportResult(
            analysis=analysis,
            entries=entries,
            invalid_entries=invalid_entries,
        )

    finally:
        workbook.close()


def _analyze_sheet(worksheet):
    """
    Detect candidate accession columns in one worksheet.
    """

    values_by_column = {}

    for row_number, row in enumerate(
        worksheet.iter_rows(
            values_only=True
        ),
        start=1,
    ):

        for column_index, value in enumerate(
            row,
            start=1,
        ):

            text = _normalize_value(
                value
            )

            if not text:
                continue

            values_by_column.setdefault(
                column_index,
                [],
            ).append(
                (
                    row_number,
                    text,
                )
            )

    available_columns = (
        _build_available_columns(
            values_by_column
        )
    )
    candidates = []

    for column_index, values in (
        values_by_column.items()
    ):

        accession_values = [
            (
                row_number,
                text,
            )
            for row_number, text
            in values
            if is_valid_assembly_accession(
                text
            )
        ]

        if not accession_values:
            continue

        first_accession_row = (
            accession_values[0][0]
        )

        label = _infer_column_label(
            values,
            first_accession_row,
            column_index,
        )

        valid_count = len(
            accession_values
        )

        non_empty_count = len(
            values
        )

        match_ratio = (
            valid_count
            / non_empty_count
        )

        examples = tuple(
            text.upper()
            for _, text
            in accession_values[:3]
        )

        score = _calculate_candidate_score(
            label=label,
            valid_count=valid_count,
            match_ratio=match_ratio,
        )

        candidates.append(
            ColumnCandidate(
                column_index=column_index,
                column_letter=(
                    get_column_letter(
                        column_index
                    )
                ),
                label=label,
                valid_accessions=(
                    valid_count
                ),
                non_empty_values=(
                    non_empty_count
                ),
                match_ratio=(
                    match_ratio
                ),
                score=score,
                examples=examples,
            )
        )

    candidates.sort(
        key=lambda candidate: (
            candidate.score,
            candidate.valid_accessions,
            candidate.match_ratio,
        ),
        reverse=True,
    )

    if not candidates:

        return SheetAnalysis(
            sheet_name=(
                worksheet.title
            ),
            available_columns=available_columns,
            candidate_columns=(),
            detected_column_index=None,
            detected_column_label=None,
            ambiguous=False,
        )

    best_candidate = candidates[0]

    ambiguous = False

    if len(candidates) > 1:

        second_candidate = (
            candidates[1]
        )

        # If two candidates receive very similar scores, the GUI
        # should later ask the user to explicitly confirm the column.
        if (
            best_candidate.score
            - second_candidate.score
            < 10
        ):
            ambiguous = True

    return SheetAnalysis(
        sheet_name=(
            worksheet.title
        ),
        available_columns=available_columns,
        candidate_columns=tuple(
            candidates
        ),
        detected_column_index=(
            best_candidate.column_index
        ),
        detected_column_label=(
            best_candidate.label
        ),
        ambiguous=ambiguous,
    )


def _calculate_candidate_score(
    label,
    valid_count,
    match_ratio,
):
    """
    Rank possible accession columns.

    Content is the main criterion. Column names are used only as a
    secondary hint to improve automatic detection.
    """

    score = 0.0

    # Primary signal: actual Assembly Accession values.
    score += valid_count

    score += (
        match_ratio
        * 20
    )

    normalized_label = (
        label.lower()
    )

    # Secondary hints only.
    if (
        "assembly"
        in normalized_label
        and "accession"
        in normalized_label
    ):
        score += 30

    elif "accession" in normalized_label:
        score += 15

    # A paired assembly column may also contain GCF/GCA values but is
    # normally not the user's primary accession column.
    if "paired" in normalized_label:
        score -= 10

    return score


def _infer_column_label(
    values,
    first_accession_row,
    column_index,
):
    """
    Use the closest non-accession text preceding the first Assembly
    Accession as the probable column label.
    """

    previous_text_values = [
        text
        for row_number, text
        in values
        if (
            row_number
            < first_accession_row
            and not is_valid_assembly_accession(
                text
            )
        )
    ]

    if previous_text_values:
        return previous_text_values[-1]

    return (
        f"Column "
        f"{get_column_letter(column_index)}"
    )


def _get_candidate_label(
    sheet_analysis,
    column_index,
):
    for candidate in (
        sheet_analysis
        .candidate_columns
    ):
        if (
            candidate.column_index
            == column_index
        ):
            return candidate.label

    for column in (
        sheet_analysis
        .available_columns
    ):
        if (
            column.column_index
            == column_index
        ):
            return column.label

    return (
        f"Column "
        f"{get_column_letter(column_index)}"
    )


def _looks_like_accession_header(
    value,
):
    normalized = (
        value.strip()
        .lower()
        .replace("_", " ")
    )

    return (
        "accession"
        in normalized
    )


def _normalize_value(value):
    if value is None:
        return ""

    return str(
        value
    ).strip()


def _validate_file(file_path):
    path = Path(
        file_path
    ).expanduser().resolve()

    if not path.exists():
        raise SpreadsheetImportError(
            f"Spreadsheet not found: {path}"
        )

    if not path.is_file():
        raise SpreadsheetImportError(
            f"Path is not a file: {path}"
        )

    if path.suffix.lower() != ".xlsx":
        raise SpreadsheetImportError(
            "Only .xlsx spreadsheets are supported "
            "at this stage."
        )

    return path

def _build_available_columns(
    values_by_column,
):
    columns = []

    for column_index, values in sorted(
        values_by_column.items()
    ):
        label = _infer_general_column_label(
            values,
            column_index,
        )

        examples = _get_column_examples(
            values,
            label,
        )

        columns.append(
            SpreadsheetColumn(
                column_index=column_index,
                column_letter=(
                    get_column_letter(
                        column_index
                    )
                ),
                label=label,
                examples=examples,
            )
        )

    return tuple(columns)


def _infer_general_column_label(
    values,
    column_index,
):
    if not values:
        return (
            f"Column "
            f"{get_column_letter(column_index)}"
        )

    first_value = values[0][1]

    # If the first value is already an accession, the spreadsheet
    # probably has no header for this column.
    if is_valid_assembly_accession(
        first_value
    ):
        return (
            f"Column "
            f"{get_column_letter(column_index)}"
        )

    return first_value


def _get_column_examples(
    values,
    label,
):
    examples = []

    for _, text in values:

        if text == label:
            continue

        examples.append(
            text
        )

        if len(examples) == 3:
            break

    return tuple(
        examples
    )